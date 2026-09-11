"""Programmatically generated PDF/DOCX fixtures for parsing tests — no
binaries committed to the repo. `reportlab`/`python-docx` are worker
dev-only dependencies.

Uses `reportlab.platypus` flowables (not raw `canvas.drawString`) so each
heading/paragraph is its own distinct PDF text object with real spacing —
Docling's layout model clusters bare absolute-positioned `drawString` lines
into one merged paragraph when they're visually close, which doesn't
exercise per-item node splitting the way a real Word/LibreOffice-exported
regulation PDF would.
"""

import io

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)

_STYLES = getSampleStyleSheet()
_TITLE = ParagraphStyle("FixtureTitle", parent=_STYLES["Title"], fontSize=16, spaceAfter=18)
_HEADING = ParagraphStyle("FixtureHeading", parent=_STYLES["Heading1"], spaceBefore=12, spaceAfter=10)
_BODY = ParagraphStyle("FixtureBody", parent=_STYLES["BodyText"], spaceBefore=24, spaceAfter=24)
_NESTED = ParagraphStyle("FixtureNested", parent=_BODY, leftIndent=1.5 * cm)


def digital_text_pdf() -> bytes:
    """A small multi-page PDF with a title, headings, a numbered list with
    nested letter items, and a table — enough to exercise Level 1 native
    extraction without any OCR escalation.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)

    story = [
        Paragraph("PERATURAN CONTOH NOMOR 1 TAHUN 2026", _TITLE),
        Paragraph("TENTANG PENGUJIAN DOKUMEN", _BODY),
        PageBreak(),
        Paragraph("BAB III", _HEADING),
        Paragraph("TAHAP PERENCANAAN", _HEADING),
        Paragraph("10. Umum", _BODY),
        Paragraph("Penjelasan umum mengenai tahap perencanaan.", _NESTED),
        Paragraph("11. Urut-urutan Kegiatan", _BODY),
        Paragraph("a. Kegiatan pertama dalam urutan.", _NESTED),
        Paragraph("b. Kegiatan kedua dalam urutan.", _NESTED),
        Paragraph("12. Dukungan", _BODY),
        Paragraph("a. Dukungan teknis dan administratif.", _NESTED),
        PageBreak(),
        Table(
            [
                ["No", "Nama", "Jumlah"],
                ["1", "Item Pertama", "10"],
                ["2", "Item Kedua", "20"],
            ],
            style=TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0)),
                    ("FONTSIZE", (0, 0), (-1, -1), 10),
                ]
            ),
        ),
    ]
    doc.build(story)
    return buffer.getvalue()


def bab_pasal_ayat_pdf() -> bytes:
    """BAB -> Pasal -> Ayat -> Huruf — the classic Indonesian regulatory
    grammar M4's ARTICLE_BASED interpreter must recognize.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)

    story = [
        Paragraph("PERATURAN CONTOH NOMOR 2 TAHUN 2026", _TITLE),
        Paragraph("TENTANG PENGUJIAN STRUKTUR PASAL", _BODY),
        PageBreak(),
        Paragraph("BAB I", _HEADING),
        Paragraph("KETENTUAN UMUM", _HEADING),
        Paragraph("Pasal 1", _BODY),
        Paragraph("Dalam Peraturan ini yang dimaksud dengan:", _NESTED),
        Paragraph("(1) Ayat pertama menjelaskan definisi umum.", _NESTED),
        Paragraph("(2) Ayat kedua menjelaskan ruang lingkup.", _NESTED),
        Paragraph("a. Huruf pertama dari ayat kedua.", _NESTED),
        Paragraph("b. Huruf kedua dari ayat kedua.", _NESTED),
        PageBreak(),
        Paragraph("BAB II", _HEADING),
        Paragraph("KETENTUAN LEBIH LANJUT", _HEADING),
        Paragraph("Pasal 2", _BODY),
        Paragraph("(1) Ayat tunggal pada pasal kedua.", _NESTED),
    ]
    doc.build(story)
    return buffer.getvalue()


def decision_pdf() -> bytes:
    """Menimbang/Mengingat/Memutuskan -> KESATU/KEDUA — the decision-preamble
    grammar M4's DECISION_BASED interpreter must recognize.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)

    story = [
        Paragraph("KEPUTUSAN DIREKTUR NOMOR 1 TAHUN 2026", _TITLE),
        Paragraph("TENTANG PENETAPAN CONTOH", _BODY),
        Paragraph("Menimbang:", _HEADING),
        Paragraph("a. bahwa diperlukan pengaturan contoh;", _NESTED),
        Paragraph("b. bahwa berdasarkan pertimbangan tersebut perlu ditetapkan keputusan;", _NESTED),
        Paragraph("Mengingat:", _HEADING),
        Paragraph("1. Undang-Undang Nomor 1 Tahun 2020;", _NESTED),
        PageBreak(),
        Paragraph("Memutuskan:", _HEADING),
        Paragraph("Menetapkan:", _BODY),
        Paragraph("KESATU : Menetapkan kebijakan contoh.", _BODY),
        Paragraph("KEDUA : Keputusan ini berlaku sejak tanggal ditetapkan.", _BODY),
    ]
    doc.build(story)
    return buffer.getvalue()


def digital_text_docx() -> bytes:
    """DOCX analogue of `digital_text_pdf` — headings, a numbered list, and a
    table — but with no page breaks at all: Docling gives DOCX no page
    provenance (`doc.pages` is empty, addendum §5), so this exercises the
    document-order (position-based) region segmentation path instead.
    """
    document = docx.Document()
    document.add_heading("PERATURAN CONTOH NOMOR 1 TAHUN 2026", level=0)
    document.add_paragraph("TENTANG PENGUJIAN DOKUMEN")
    document.add_heading("BAB III", level=1)
    document.add_heading("TAHAP PERENCANAAN", level=1)
    document.add_paragraph("10. Umum")
    document.add_paragraph("Penjelasan umum mengenai tahap perencanaan.")
    document.add_paragraph("11. Urut-urutan Kegiatan")
    document.add_paragraph("a. Kegiatan pertama dalam urutan.")
    document.add_paragraph("b. Kegiatan kedua dalam urutan.")
    table = document.add_table(rows=3, cols=3)
    rows = [
        ["No", "Nama", "Jumlah"],
        ["1", "Item Pertama", "10"],
        ["2", "Item Kedua", "20"],
    ]
    for row_idx, row_values in enumerate(rows):
        for col_idx, value in enumerate(row_values):
            table.cell(row_idx, col_idx).text = value
    document.add_heading("LAMPIRAN", level=1)
    document.add_paragraph("Lampiran ini berisi contoh formulir.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def image_only_pdf() -> bytes:
    """A PDF with no extractable text layer at all (a solid rectangle drawn
    with no text operators) — exercises the "never fail on empty text layer"
    path (addendum §14) and OCR/UNKNOWN_BLOCK handling.
    """
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    c.setFillColorRGB(0.5, 0.5, 0.5)
    c.rect(72, 400, 400, 300, fill=1, stroke=0)
    c.showPage()
    c.save()
    return buffer.getvalue()
